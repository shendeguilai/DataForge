import tempfile
import unittest
from pathlib import Path

from docx import Document

from compatible_converter import convert_compatible
from compatible_parser import has_lossless_features, parse_document, scan_blocks, split_sections
from compatible_studio_core import analyze_markdown


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "CSP_初赛模板.docx"
P7224 = ROOT / "fixtures" / "P7224C_2026_CSP-S.md"


class LosslessCompatibilityTest(unittest.TestCase):
    def test_fenced_standard_metadata_is_not_an_unrendered_question(self):
        source = (ROOT.parent.parent / 'src/main/resources/static/csp-paper-studio-template.md').read_text()
        _, header, rest = source.split('---', 2)
        for opener, closer in [('```', '```'), ('```yaml', '```'), ('~~~yml', '~~~')]:
            with self.subTest(opener=opener):
                markdown = opener + header + closer + rest
                model = parse_document(markdown)
                self.assertEqual('standard-v1', model.detected_format)
                self.assertEqual(20, len(model.modules))
                self.assertFalse(model.errors)
                self.assertEqual(0, model.dropped_blocks)
                self.assertEqual('S', model.metadata['group'])
                self.assertTrue(any(d.code == 'FENCED_METADATA' for d in model.warnings))
                self.assertTrue(model.preamble.startswith(opener))
                for module in model.modules:
                    for block in module.all_rendered_blocks():
                        if block.fenced:
                            self.assertTrue(markdown.splitlines()[block.source_start_line - 1].strip().startswith(('```', '~~~')))

    def test_metadata_detection_does_not_hide_real_preamble_code(self):
        from compatible_parser import parse_front_matter
        for content in ['title: sample', 'csp_format: 1\nprint(123)',
                        'csp_format: 1\ncustom: payload', 'csp_format: 1\ngroup: J\ngroup: S']:
            source = '```yaml\n' + content + '\n```\n## 第 1 题\n题干'
            self.assertEqual(({}, 'legacy-v04', 0), parse_front_matter(source))
            self.assertTrue(any(d.code == 'UNRENDERED_BLOCK' and d.line == 1 for d in parse_document(source).errors))
        self.assertEqual(({}, 'legacy-v04', 0), parse_front_matter('```cpp\ncsp_format: 1\n```'))

    def test_reported_xor_reading_preserves_inputs_and_question_order_in_word(self):
        source = P7224.read_text(encoding="utf-8")
        model = parse_document(source)
        reading = next(x for x in model.modules if x.number == 17)
        inputs = ["1\n4\n-1 -1 -1 -1\n2 3 5 7", "1\n4\n-1 34 367 -1\n3178 -1 -1 3333"]
        for item, expected in zip(reading.choices, inputs):
            self.assertEqual(expected, next(b.content for b in item.blocks if b.fenced))
        self.assertEqual(3, len(reading.choices))
        self.assertFalse(model.errors)
        with tempfile.TemporaryDirectory(prefix="csp_xor_regression_") as directory:
            output = Path(directory) / "reading.docx"
            convert_compatible(P7224, output, TEMPLATE, "auto")
            document = Document(output)
            # Read paragraphs and tables in actual document order; checking
            # table presence alone misses misplaced input blocks.
            ordered = ["".join(node.xpath(".//w:t/text()")) for node in document.element.body]
            first = ordered.index("14-1 -1 -1 -12 3 5 7")
            second = ordered.index("14-1 34 367 -13178 -1 -1 3333")
            self.assertIn("输入为", ordered[first - 1])
            self.assertIn("程序输出为", ordered[first + 1])
            self.assertIn("1 6 2 7", ordered[first + 2])
            self.assertIn("输入为", ordered[second - 1])
            self.assertIn("程序输出的第三个数", ordered[second + 1])
            self.assertIn("3333", ordered[second + 2])
            self.assertFalse(any("```" in text for text in ordered))

    def test_reported_broken_fences_cannot_silently_shift_choice_options(self):
        source = P7224.read_text(encoding="utf-8")
        source = source.replace(
            "```text\n1\n4\n-1 -1 -1 -1\n2 3 5 7\n```",
            "```text\n  ```\n\n1\n4\n-1 -1 -1 -1\n2 3 5 7\n````yaml",
        ).replace("3178 -1 -1 3333\n```", "3178 -1 -1 3333\n````")
        model = parse_document(source)
        self.assertTrue(any(x.code == "EMPTY_FENCE" and x.line > 1 for x in model.warnings))
        self.assertTrue(any(x.code == "SUBQUESTION_MISMATCH" for x in model.errors))
        self.assertGreater(analyze_markdown(source)["counts"]["error"], 0)
        with tempfile.TemporaryDirectory(prefix="csp_broken_fences_") as directory:
            path = Path(directory) / "broken.md"
            output = Path(directory) / "broken.docx"
            path.write_text(source, encoding="utf-8")
            from converter import ConvertError
            with self.assertRaisesRegex(ConvertError, "第 .* 行"):
                convert_compatible(path, output, TEMPLATE)
            self.assertFalse(output.exists())

    def test_typed_fences_and_nested_backticks_are_lossless(self):
        diagnostics = []
        blocks = scan_blocks(
            """前言\n\n   ~~~text   \n a`b\n  c\n   ~~~\n\n```input\n7\n8 1 6\n```\n\n```math\nx_i^2\n```\n""",
            diagnostics=diagnostics,
        )
        fenced = [x for x in blocks if x.fenced]
        self.assertEqual(["preformatted", "preformatted", "math"], [x.kind for x in fenced])
        self.assertEqual("a`b\nc", fenced[0].content)
        self.assertEqual("7\n8 1 6", fenced[1].content)
        self.assertFalse(any(x.code == "UNCLOSED_FENCE" for x in diagnostics))

    def test_crlf_tilde_and_deeply_indented_fences(self):
        diagnostics = []
        blocks = scan_blocks(
            "1. 示例\r\n        ~~~input  \r\n        7\r\n        1. data\r\n        ~~~\r\n2. 下一项\r\n",
            diagnostics=diagnostics,
        )
        fenced = [x for x in blocks if x.fenced]
        self.assertEqual(1, len(fenced))
        self.assertEqual("input", fenced[0].language)
        self.assertEqual("7\n1. data", fenced[0].content)
        self.assertFalse(diagnostics)

    def test_section_heading_inside_fence_is_not_a_question_boundary(self):
        preamble, sections = split_sections(
            "## 第 1 题\n"
            "```text\n"
            "## 第 99 题\n"
            "~~~ not a heading\n"
            "```\n"
            "## 第 2 题\n"
        )
        self.assertEqual("", preamble)
        self.assertEqual([1, 2], [x["number"] for x in sections])

    def test_indented_question_headings_are_supported(self):
        preamble, sections = split_sections(
            "\ufeff  ## 第 1 题\n题干\n  ## 第 2 题\n题干\n"
        )
        self.assertEqual("", preamble)
        self.assertEqual([1, 2], [x["number"] for x in sections])

    def test_plaintext_alias_is_preformatted_data(self):
        blocks = scan_blocks("```plaintext\na  b\n```")
        self.assertEqual("preformatted", blocks[0].kind)
        self.assertEqual("data", blocks[0].role)

    def test_bom_prefixed_fence_is_scanned(self):
        blocks = scan_blocks("\ufeff```text\n保留 BOM 后的内容\n```")
        self.assertEqual("preformatted", blocks[0].kind)
        self.assertEqual("保留 BOM 后的内容", blocks[0].content)

    def test_full_width_option_separator_uses_lossless_path(self):
        markdown = "\n".join([
            "## 第 1 题", "题干", "- A．甲", "- B．乙", "- C．丙", "- D．丁",
        ])
        self.assertTrue(has_lossless_features(markdown))
        model = parse_document(markdown)
        self.assertEqual(["A. 甲", "B. 乙", "C. 丙", "D. 丁"], model.modules[0].options)
        self.assertFalse(model.modules[0].errors)

    def test_p7224_is_fully_covered_when_available(self):
        if not P7224.is_file():
            self.skipTest("用户示例不在当前机器")
        model = parse_document(P7224.read_text(encoding="utf-8-sig"))
        self.assertEqual("legacy-v04", model.detected_format)
        self.assertEqual(20, len(model.modules))
        self.assertEqual(0, model.dropped_blocks)
        self.assertFalse(model.modules[0].errors)
        self.assertEqual(15, sum(x.kind == "choice" for x in model.modules))
        self.assertEqual(3, sum(x.kind == "reading" for x in model.modules))
        self.assertEqual(2, sum(x.kind == "fill" for x in model.modules))
        q16 = next(x for x in model.modules if x.number == 16)
        self.assertIn("7\n8 1 6 3 5 7 2", q16.choices[0].blocks[1].content)
        q19 = next(x for x in model.modules if x.number == 19)
        self.assertEqual("x1 ≤ x2 ≤ ... ≤ xk\ny1 ≤ y2 ≤ ... ≤ yk", q19.prose_blocks[1].content)
        analysis = analyze_markdown(P7224.read_text(encoding="utf-8-sig"))
        self.assertEqual(2, analysis["schema_version"])
        self.assertIn("csp_format: 1", analysis["migration_hint"])

        with tempfile.TemporaryDirectory(prefix="csp_compat_test_") as directory:
            output = Path(directory) / "p7224.docx"
            convert_compatible(P7224, output, TEMPLATE, "auto")
            document = Document(output)
            all_text = "\n".join(x.text for x in document.paragraphs)
            self.assertNotIn("text (A+B)", all_text)
            self.assertGreaterEqual(sum(1 for table in document.tables if len(table.columns) == 1), 10)
            self.assertTrue(any("x1 ≤ x2" in table.rows[0].cells[0].text for table in document.tables if len(table.columns) == 1))

    def test_p7224_numbering_modes(self):
        for mode, expected_prefix in (("auto", "(1)."), ("local", "(1)."), ("global", "16.")):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory(prefix="csp_numbering_") as directory:
                output = Path(directory) / "numbering.docx"
                convert_compatible(P7224, output, TEMPLATE, mode)
                paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
                self.assertTrue(any(text.startswith(expected_prefix) for text in paragraphs))

    def test_standard_template_has_twenty_modules(self):
        template = ROOT.parent.parent / "src/main/resources/static/csp-paper-studio-template.md"
        model = parse_document(template.read_text(encoding="utf-8"))
        self.assertEqual("standard-v1", model.detected_format)
        self.assertEqual(list(range(1, 21)), [x.number for x in model.modules])
        self.assertFalse(model.errors)
        self.assertEqual(0, model.dropped_blocks)

    def test_unclosed_and_unknown_fences_are_diagnosed(self):
        unknown = """---\ncsp_format: 1\n---\n## 第 1 题\n### 题干\n题干\n```weird\na`b\n```\n### 选项\n- A. a\n- B. b\n- C. c\n- D. d\n"""
        model = parse_document(unknown)
        self.assertTrue(any(x.code == "UNKNOWN_FENCE" for x in model.diagnostics))
        self.assertEqual(0, model.dropped_blocks)
        blocks = scan_blocks("## 第 1 题\n```text\nnever closed\n", diagnostics=[])
        self.assertEqual("preformatted", blocks[-1].kind)

    def test_program_aliases_and_tilde_fences_select_lossless_renderer(self):
        for opener, closer in (("```c++", "```"), ("```cxx", "```"), ("~~~cpp", "~~~")):
            markdown = "\n".join([
                "## 第 1 题", "题干", opener, "int main() {}", closer,
                "本题共 2 分", "- A. a", "- B. b", "- C. c", "- D. d",
            ])
            self.assertTrue(has_lossless_features(markdown), opener)
            model = parse_document(markdown)
            self.assertFalse(model.modules[0].errors, opener)
            self.assertEqual("code", next(x for x in model.modules[0].blocks if x.fenced).kind)

    def test_headings_inside_fences_do_not_split_standard_fields(self):
        markdown = "\n".join([
            "---", "csp_format: 1", "year: 2026", "group: S", "---",
            "## 第 1 题", "### 类型：单项选择题", "### 题干",
            "```text", "### 选项", "literal heading", "```", "真正的题干。",
            "### 选项", "- A. a", "- B. b", "- C. c", "- D. d",
        ])
        model = parse_document(markdown)
        self.assertFalse(model.modules[0].errors)
        self.assertEqual(["A. a", "B. b", "C. c", "D. d"], model.modules[0].options)
        self.assertIn("### 选项\nliteral heading", next(x for x in model.modules[0].blocks if x.fenced).content)

    def test_lossless_renderer_keeps_math_code_and_long_code_rows(self):
        """Exercise the Word-only parts of the standard (format 1) path.

        The legacy converter has its own regression suite.  This case makes
        sure the compatibility renderer does not accidentally feed dollars in
        a C++ string to Pandoc, and that dynamically extended code tables keep
        each source line together when Word paginates them.
        """
        source = (ROOT.parent.parent / "src/main/resources/static/csp-paper-studio-template.md").read_text(
            encoding="utf-8"
        )
        long_code = "\n".join(
            f"{index:02d}     std::cout << \"$x_i$ 2^60\"; // {'x' * 120}"
            for index in range(1, 45)
        )
        replacement = (
            "这里有行内公式 $a_i^2$，还有显示公式 \\[x^2+y^2=z^2\\]。\n\n"
            "```cpp\n"
            + long_code
            + "\n```\n\n"
            "```input\n"
            "first  1\n"
            "  second 2\n"
            "\n"
            "third\n"
            "```"
        )
        source = source.replace("在这里填写第 1 题题干。", replacement)

        model = parse_document(source)
        self.assertFalse(model.errors)
        self.assertEqual(2, next(x for x in model.modules if x.number == 1).formula_count)

        with tempfile.TemporaryDirectory(prefix="csp_word_compat_") as directory:
            markdown = Path(directory) / "standard.md"
            output = Path(directory) / "standard.docx"
            markdown.write_text(source, encoding="utf-8")
            convert_compatible(markdown, output, TEMPLATE, "auto")
            document = Document(output)

            # The first two-column table is the injected 44-line C++ block.
            code_tables = [
                table
                for table in document.tables
                if len(table.columns) == 2 and table.rows and table.rows[0].cells[0].text.strip() == "01"
            ]
            self.assertTrue(code_tables)
            code_table = code_tables[0]
            self.assertEqual(44, len(code_table.rows))
            self.assertTrue(all(row._tr.xpath("./w:trPr/w:cantSplit") for row in code_table.rows))
            self.assertIn('$x_i$ 2^60', code_table.cell(0, 1).text)

            preformatted = [table for table in document.tables if len(table.columns) == 1]
            self.assertTrue(any(len(table.rows) == 4 for table in preformatted))
            input_table = next(table for table in preformatted if len(table.rows) == 4)
            self.assertEqual("first  1", input_table.cell(0, 0).text)
            self.assertEqual("  second 2", input_table.cell(1, 0).text)
            self.assertEqual("", input_table.cell(2, 0).text.strip())
            self.assertEqual("third", input_table.cell(3, 0).text)

            # The two explicitly marked formulas are OMML; the dollar signs
            # inside the C++ string are present only as literal code text.
            self.assertEqual(2, sum(len(paragraph._p.xpath(".//m:oMath")) for paragraph in document.paragraphs))

    def test_option_detection_ignores_labels_inside_fenced_code_and_input(self):
        # A-D-looking lines in a program/data block must never become the
        # question's first options or split a nested reading item.
        markdown = """## 第 1 题
### 题干
```cpp
A. this is source code
B. this is source code
C. this is source code
D. this is source code
```
### 选项
- A. real option A
- B. real option B
- C. real option C
- D. real option D
        """
        model = parse_document(markdown)
        self.assertFalse(any(x.code in {"PARSE_ERROR", "UNRENDERED_BLOCK"} for x in model.errors))
        self.assertEqual(
            ["A. real option A", "B. real option B", "C. real option C", "D. real option D"],
            model.modules[0].options,
        )

        reading = """## 第 16 题
### 类型：阅读程序
### 程序
```cpp
int main() { return 0; }
```
### 说明
说明。
### 判断题
1. 判断。
### 单选题
1. 输入为：
   ```input
   1. data line
   - A. data line
   ```
   - A. A
   - B. B
   - C. C
   - D. D
        """
        reading_model = parse_document(reading)
        self.assertEqual([], reading_model.modules[0].errors)
        self.assertEqual(1, len(reading_model.modules[0].choices))
        self.assertIn("1. data line", reading_model.modules[0].choices[0].blocks[1].content)
        self.assertEqual("A. A", reading_model.modules[0].choices[0].options[0])

    def test_explicit_math_and_inline_code_are_not_reported_as_bare_math(self):
        markdown = """## 第 1 题
### 题干
`a_i` and $a = 2^60$ and `price=$5$`; bare 2^60 remains prose.
### 选项
- A. a
- B. b
- C. c
- D. d
        """
        model = parse_document(markdown)
        bare = [x for x in model.diagnostics if x.code == "BARE_MATH"]
        self.assertEqual(1, len(bare))
        self.assertIn("2^60", bare[0].message)

    def test_unclosed_explicit_math_is_kept_and_reported_with_line(self):
        markdown = "\n".join([
            "## 第 1 题", "题干", "第一行", "$a_i^2", "### 选项",
            "- A. a", "- B. b", "- C. c", "- D. d",
        ])
        model = parse_document(markdown)
        errors = [x for x in model.diagnostics if x.code == "MATH_PARSE_ERROR"]
        self.assertEqual(1, len(errors))
        self.assertEqual(4, errors[0].line)
        stem = "\n".join(x.content for x in model.modules[0].blocks)
        self.assertIn("$a_i^2", stem)

    def test_multiline_options_keep_wrapped_text_and_crlf(self):
        markdown = "\r\n".join([
            "---", "csp_format: 1", "year: 2026", "group: S", "---",
            "## 第 1 题", "### 类型：单项选择题", "### 题干", "题干。",
            "### 选项", "- A. 第一行", "  第二行", "", "  第三行",
            "- B. 选项二", "- C. 选项三", "- D. 选项四", "",
        ])
        model = parse_document("\ufeff" + markdown)
        self.assertEqual("standard-v1", model.detected_format)
        self.assertFalse(model.modules[0].errors)
        self.assertEqual("A. 第一行\n第二行\n\n第三行", model.modules[0].options[0])
        self.assertEqual(0, model.dropped_blocks)

        plus_bullets = markdown.replace("- A.", "+ A.").replace("- B.", "+ B.").replace("- C.", "+ C.").replace("- D.", "+ D.")
        plus_model = parse_document(plus_bullets)
        self.assertEqual(["A. 第一行\n第二行\n\n第三行", "B. 选项二", "C. 选项三", "D. 选项四"], plus_model.modules[0].options)

    def test_adjacent_mixed_fences_and_inner_backticks(self):
        diagnostics = []
        blocks = scan_blocks(
            "\n".join([
                "## 第 1 题", "```text", "A. input", "```",
                "~~~output", "B. output", "~~~",
                "````cpp", "const char *s = \"```\";", "````",
                "```weird", "literal", "```",
            ]),
            diagnostics=diagnostics,
        )
        fenced = [x for x in blocks if x.fenced]
        self.assertEqual(["preformatted", "preformatted", "code", "preformatted"], [x.kind for x in fenced])
        self.assertEqual('const char *s = "```";', fenced[2].content)
        self.assertEqual(1, len([x for x in diagnostics if x.code == "UNKNOWN_FENCE"]))
        self.assertFalse(any(x.code == "UNCLOSED_FENCE" for x in diagnostics))


if __name__ == "__main__":
    unittest.main()
