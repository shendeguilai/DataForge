#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word renderer for the lossless CSP Markdown model.

All legacy documents that only contain the V0.4 ``cpp`` fences continue to
use :func:`converter.convert`, preserving the byte-for-byte regression
baseline.  The renderer below is selected for format 1 and for documents
that contain typed input/output/math/unknown fences.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import converter as legacy
from compatible_parser import (
    formulas_in_blocks,
    has_lossless_features,
    parse_document,
    strip_math_delimiters,
)
from document_model import Block, Diagnostic, DocumentModel, Module


ConvertError = legacy.ConvertError


def _safe_fill_inline(paragraph, text: str, bank: dict):
    """Fill a paragraph while retaining invalid formulas as literal text."""
    legacy.clear_paragraph(paragraph)
    value = str(text).replace("\\(", "$").replace("\\)", "$")
    token_re = re.compile(r"(\$\$.*?\$\$|\\\[.*?\\\]|(?<!\\)\$[^$\n]*?\$|\\\([^\n]*?\\\)|`[^`\n]*`)", re.S)
    pos = 0
    for match in token_re.finditer(value):
        if match.start() > pos:
            legacy.add_text_run(paragraph, value[pos:match.start()])
        token = match.group(0)
        if token.startswith("`"):
            legacy.add_text_run(paragraph, token[1:-1], code=True)
        else:
            if token.startswith("$$"):
                raw = token[2:-2].strip()
            elif token.startswith("\\["):
                raw = token[2:-2].strip()
            elif token.startswith("\\("):
                raw = token[2:-2].strip()
            else:
                raw = token[1:-1].strip()
            raw = legacy.normalize_math_tex(raw)
            if raw in bank:
                legacy.add_math(paragraph, raw, bank)
            else:
                # Pandoc can reject an individual TeX expression.  The source
                # remains visible and can be corrected without losing text.
                legacy.add_text_run(paragraph, token)
        pos = match.end()
    if pos < len(value):
        legacy.add_text_run(paragraph, value[pos:])


def _add_text_paragraph(doc, style, text, bank, keep_next=False):
    paragraph = doc.add_paragraph(style=style)
    _safe_fill_inline(paragraph, text, bank)
    if keep_next:
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "C8CDD2")


def _set_cell_shading(cell, fill="F8FAFC"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_preformatted_table(doc, block: Block):
    """Render input/output/data as a one-column, whitespace-preserving table."""
    lines = block.content.split("\n") or [""]
    table = doc.add_table(rows=max(1, len(lines)), cols=1)
    table.autofit = False
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for row, line in zip(table.rows, lines):
        _set_cant_split(row)
        cell = row.cells[0]
        cell.width = Inches(6.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_border(cell)
        _set_cell_shading(cell)
        paragraph = cell.paragraphs[0]
        try:
            paragraph.style = "CSPJ 代码"
        except Exception:
            pass
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        legacy.clear_paragraph(paragraph)
        run = paragraph.add_run(line if line else " ")
        legacy.set_run_font(run, "Consolas", "Noto Sans Mono CJK SC", 9.0)
    return table


def add_compat_code_table(doc, refs, code):
    """Render a V0.4 code table and make every source line indivisible.

    Word is allowed to split a table row across pages unless ``w:cantSplit``
    is present.  That is especially noticeable for a long line at the bottom
    of a page: the line number can remain on one page while the code text is
    moved to the next.  The V0.4 renderer remains untouched; the compatibility
    renderer opts into the safer row setting for its lossless path only.
    """
    table = legacy.add_code_table(doc, refs, code)
    for row in table.rows:
        _set_cant_split(row)
    return table


def add_compat_option_table(doc, refs, options, bank):
    """V0.4 option table with the safe inline formula fallback."""
    shape = legacy.choose_option_shape(options)
    table = legacy.clone_table_to_doc(doc, refs[shape])
    legacy._trim_unused_option_rows(table, shape, len(options))
    cells = legacy.flatten_cells(table)
    if len(cells) < len(options):
        raise ConvertError(f"选项表容量不足：{shape}")
    for index, value in enumerate(options):
        cell = cells[index]
        paragraph = cell.paragraphs[0]
        try:
            paragraph.style = "CSPJ 选项"
        except Exception:
            pass
        _safe_fill_inline(paragraph, value, bank)
        for extra in cell.paragraphs[1:]:
            extra._element.getparent().remove(extra._element)
    for cell in cells[len(options):]:
        paragraph = cell.paragraphs[0]
        legacy.clear_paragraph(paragraph)
    return table


def _render_block_sequence(doc, refs, blocks: list[Block], bank: dict, *, prefix="", paragraph_style="CSPJ 题目"):
    if not blocks:
        if prefix:
            _add_text_paragraph(doc, paragraph_style, prefix, bank, keep_next=True)
        return
    first_text = True
    for block in blocks:
        if block.kind == "paragraph":
            text = (prefix if first_text else "") + block.content
            _add_text_paragraph(doc, paragraph_style, text, bank, keep_next=True)
            first_text = False
        elif block.kind == "code":
            if first_text and prefix:
                _add_text_paragraph(doc, paragraph_style, prefix, bank, keep_next=True)
                first_text = False
            add_compat_code_table(doc, refs, block.content)
        elif block.kind == "preformatted":
            if first_text and prefix:
                _add_text_paragraph(doc, paragraph_style, prefix, bank, keep_next=True)
                first_text = False
            add_preformatted_table(doc, block)
        elif block.kind == "math":
            # Math fences normally contain raw TeX, but accepting one pair of
            # pasted delimiters keeps `````math````/``$$...$$`` interchangeable.
            # Display wrapping also preserves newlines in multi-line formulas.
            value = f"$${strip_math_delimiters(block.content)}$$"
            _add_text_paragraph(doc, paragraph_style, (prefix if first_text else "") + value, bank, keep_next=True)
            first_text = False


def _safe_math_bank(model: DocumentModel):
    option_blocks = []
    for module in model.modules:
        values = list(module.options)
        for item in [*module.judges, *module.choices, *module.groups]:
            values.extend(item.options)
        option_blocks.extend(Block("paragraph", "prose", content=value) for value in values)
    formulas = formulas_in_blocks([b for m in model.modules for b in m.all_rendered_blocks()] + option_blocks)
    if not formulas:
        return {}
    try:
        return legacy.build_math_bank(formulas)
    except ConvertError:
        # Retry one formula at a time so a malformed expression cannot erase
        # otherwise valid equations.  The caller receives line diagnostics.
        bank = {}
        for formula in formulas:
            try:
                bank.update(legacy.build_math_bank([formula]))
            except ConvertError:
                block = next((b for m in model.modules for b in m.all_rendered_blocks() if formula in b.content), None)
                line = block.source_start_line if block else 1
                model.diagnostics.append(Diagnostic("warning", "MATH_PARSE_ERROR", line, f"公式无法转换，已保留原文：${formula}$"))
        return bank


def _heading_or_default(body: str, prefix: str, default: str):
    match = re.search(rf"(?m)^#{1,6}\s*({re.escape(prefix)}[^\n]*)$", body)
    return match.group(1).strip() if match else default


def render_model(model: DocumentModel, output: Path, template: Path, numbering="auto"):
    if model.errors or model.dropped_blocks:
        details = "；".join(f"第 {x.line} 行：{x.message}" for x in model.errors[:5])
        raise ConvertError(details or "Markdown 中存在无法渲染的区块。")
    bank = _safe_math_bank(model)
    template_doc = Document(template)
    if len(template_doc.tables) < 36:
        raise ConvertError("模板不符合要求。")
    refs = legacy.setup_refs(template_doc)
    doc = Document(template)
    legacy.clear_document_body(doc)
    year = model.metadata.get("year", "20XX")
    group = model.metadata.get("group", "J")
    level = model.metadata.get("level", "提高级" if group == "S" else "入门级")
    global_numbering = group == "J" if numbering == "auto" else numbering == "global"

    _add_text_paragraph(doc, "CSPJ 标题", f"{year} CCF 非专业级别软件能力认证第一轮", bank)
    _add_text_paragraph(doc, "CSPJ 副标题", f"（CSP-{group}1）{level} C++ 语言试题", bank)
    _add_text_paragraph(doc, "CSPJ 考试信息", "考试时间：120 分钟    试卷满分：100 分", bank)
    _add_text_paragraph(doc, "Normal", "", bank)
    _add_text_paragraph(doc, "CSPJ 注意标题", "考生注意事项：", bank, keep_next=True)
    for note in [
        "• 试题分为单项选择题、阅读程序题和完善程序题三部分。",
        "• 所有程序均使用 C++ 编写；除特殊说明外，程序输入不超过数组或字符串定义的范围。",
        "• 单项选择题和完善程序题每题有且仅有一个正确选项。",
        "• 阅读程序中的判断题，正确选“A”，错误选“B”；其余为单项选择题。" if group == "S" else "• 阅读程序中的判断题，正确填“√”，错误填“×”。",
    ]:
        _add_text_paragraph(doc, "CSPJ 注意事项", note, bank)

    _add_text_paragraph(doc, "CSPJ 大题标题", "一、单项选择题（共 15 题，每题 2 分，共计 30 分；每题有且仅有一个正确选项）", bank, keep_next=True)
    modules = {x.number: x for x in model.modules}
    for number in range(1, 16):
        module = modules.get(number)
        if module is None:
            continue
        _render_block_sequence(doc, refs, module.blocks, bank, prefix=f"{number}. ", paragraph_style="CSPJ 题目")
        add_compat_option_table(doc, refs, module.options, bank)

    read_title = _heading_or_default(modules.get(16, Module(16, "reading", "", 1)).body, "二、", "二、阅读程序")
    _add_text_paragraph(doc, "CSPJ 大题标题", read_title, bank, keep_next=True)
    current_no = 16
    for index, number in enumerate((16, 17, 18), 1):
        module = modules.get(number)
        if module is None or module.code is None:
            continue
        _add_text_paragraph(doc, "CSPJ 小节标题", f"（{index}） 阅读下列程序，回答问题。", bank, keep_next=True)
        add_compat_code_table(doc, refs, module.code.content)
        _render_block_sequence(doc, refs, module.note_blocks, bank, paragraph_style="CSPJ 正文")
        _add_text_paragraph(doc, "CSPJ 类型标题", "·判断题：", bank, keep_next=True)
        for j, item in enumerate(module.judges, 1):
            if global_numbering:
                label = f"{current_no}. "; current_no += 1
            else:
                label = f"({j}). "
            text = item.text
            if group == "J" and not re.search(r"[（(]\s*[）)]\s*$", text):
                text += "（ ）"
            item_blocks = list(item.blocks)
            if group == "J" and item_blocks and item_blocks[0].kind == "paragraph" and text != item.text:
                item_blocks[0] = copy.copy(item_blocks[0])
                item_blocks[0].content = item_blocks[0].content + "（ ）"
            _render_block_sequence(doc, refs, item_blocks, bank, prefix=label, paragraph_style="CSPJ 题目")
            if group != "J":
                add_compat_option_table(doc, refs, ["A. 正确", "B. 错误"], bank)
        _add_text_paragraph(doc, "CSPJ 类型标题", "·单选题：", bank, keep_next=True)
        for k, item in enumerate(module.choices, 1):
            if global_numbering:
                label = f"{current_no}. "; current_no += 1
            else:
                label = f"({len(module.judges) + k}). "
            _render_block_sequence(doc, refs, item.blocks, bank, prefix=label, paragraph_style="CSPJ 题目")
            add_compat_option_table(doc, refs, item.options, bank)

    fill_title = _heading_or_default(modules.get(19, Module(19, "fill", "", 1)).body, "三、", "三、完善程序（单选题，每小题 3 分，共计 30 分）")
    _add_text_paragraph(doc, "CSPJ 大题标题", fill_title, bank, keep_next=True)
    for index, number in enumerate((19, 20), 1):
        module = modules.get(number)
        if module is None or module.code is None:
            continue
        prose = module.prose_blocks
        if prose:
            _render_block_sequence(doc, refs, prose, bank, prefix=f"（{index}）", paragraph_style="CSPJ 正文")
        else:
            _add_text_paragraph(doc, "CSPJ 小节标题", f"（{index}）", bank, keep_next=True)
        add_compat_code_table(doc, refs, module.code.content)
        for sub, item in enumerate(module.groups, 1):
            if global_numbering:
                label = f"{current_no}. "; current_no += 1
            else:
                label = f"({sub}). "
            _render_block_sequence(doc, refs, item.blocks, bank, prefix=label, paragraph_style="CSPJ 题目")
            add_compat_option_table(doc, refs, item.options, bank)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def convert_compatible(src: Path, output: Path, template: Path, numbering="auto"):
    markdown = src.read_text(encoding="utf-8-sig")
    if not has_lossless_features(markdown):
        return legacy.convert(src, output, template, numbering)
    model = parse_document(markdown)
    render_model(model, output, template, numbering)
